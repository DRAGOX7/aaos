"""Deterministic structural extraction from DOCX bytes."""

from __future__ import annotations

from io import BytesIO
from zipfile import BadZipFile, is_zipfile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml.etree import LxmlError  # type: ignore[import-untyped]

from aaos.documents.contracts import (
    DocumentIngestionFailure,
    ExtractedStructureUnit,
    IngestionErrorCode,
    ParsedDocx,
    StructureUnitKind,
)


class DocxProcessor:
    """Extract bounded structural provenance without inventing page semantics."""

    def process(self, body: bytes) -> ParsedDocx:
        if not isinstance(body, bytes) or not body:
            raise DocumentIngestionFailure(
                IngestionErrorCode.INVALID_DOCX
            )

        try:
            archive_like = is_zipfile(BytesIO(body))
        except (OSError, ValueError):
            archive_like = False

        if not archive_like:
            raise DocumentIngestionFailure(
                IngestionErrorCode.INVALID_DOCX
            )

        try:
            document = Document(BytesIO(body))
        except (
            BadZipFile,
            KeyError,
            LxmlError,
            PackageNotFoundError,
            ValueError,
        ):
            raise DocumentIngestionFailure(
                IngestionErrorCode.INVALID_DOCX
            ) from None

        units: list[ExtractedStructureUnit] = []
        ordinal = 1
        table_index = 0

        try:
            for block in document.iter_inner_content():
                if isinstance(block, Paragraph):
                    heading_level = _heading_level(block)
                    units.append(
                        ExtractedStructureUnit(
                            ordinal=ordinal,
                            kind=(
                                StructureUnitKind.HEADING
                                if heading_level is not None
                                else StructureUnitKind.PARAGRAPH
                            ),
                            text=block.text,
                            heading_level=heading_level,
                            table_index=None,
                            row_index=None,
                            column_index=None,
                        )
                    )
                    ordinal += 1
                    continue

                if not isinstance(block, Table):
                    raise DocumentIngestionFailure(
                        IngestionErrorCode.DOCX_PARSE_FAILED
                    )

                seen_physical_cells: list[object] = []

                for row_index, row in enumerate(block.rows):
                    column_index = row.grid_cols_before

                    for cell in row.cells:
                        physical_cell = cell._tc

                        if (
                            cell.grid_span != 1
                            or any(
                                physical_cell is existing
                                for existing in seen_physical_cells
                            )
                        ):
                            raise DocumentIngestionFailure(
                                IngestionErrorCode.UNSUPPORTED_DOCX_STRUCTURE
                            )

                        seen_physical_cells.append(physical_cell)

                        if cell.tables:
                            raise DocumentIngestionFailure(
                                IngestionErrorCode.UNSUPPORTED_DOCX_STRUCTURE
                            )

                        units.append(
                            ExtractedStructureUnit(
                                ordinal=ordinal,
                                kind=StructureUnitKind.TABLE_CELL,
                                text="\n".join(
                                    paragraph.text
                                    for paragraph in cell.paragraphs
                                ),
                                heading_level=None,
                                table_index=table_index,
                                row_index=row_index,
                                column_index=column_index,
                            )
                        )
                        ordinal += 1
                        column_index += 1

                table_index += 1

        except DocumentIngestionFailure:
            raise
        except (
            AttributeError,
            IndexError,
            KeyError,
            LxmlError,
            TypeError,
            ValueError,
        ):
            raise DocumentIngestionFailure(
                IngestionErrorCode.DOCX_PARSE_FAILED
            ) from None

        return ParsedDocx(units=tuple(units))


def _heading_level(paragraph: Paragraph) -> int | None:
    style = paragraph.style
    style_name = style.name if style is not None else ""

    prefix = "Heading "
    if not style_name.startswith(prefix):
        return None

    suffix = style_name[len(prefix) :]
    if len(suffix) != 1 or suffix not in "123456789":
        return None

    return int(suffix)
